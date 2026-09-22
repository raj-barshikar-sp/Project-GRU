"""Turns a DeckSpec into a .pptx file on the SailPoint master.

The agent decides what the slides say. This code opens the company template
and fills its placeholders. Models cannot do that part: a .pptx is a zip of XML.
"""

from __future__ import annotations

import io
import os
import tempfile
import zipfile
from pathlib import Path

from pptx import Presentation
from pptx.enum.shapes import PP_PLACEHOLDER
from pptx.oxml.ns import qn
from pptx.util import Emu

from ..contracts import DeckSpec, Slide, SlideLayout

OUTPUT_DIR = Path("output/decks")
TEMPLATE_FILENAME = "SailPoint-PowerPointTemplate-202600707.potx"
TEMPLATE_CT = (
    "application/vnd.openxmlformats-officedocument.presentationml.template.main+xml"
)
PRESENTATION_CT = (
    "application/vnd.openxmlformats-officedocument.presentationml.presentation.main+xml"
)

# Names from the SailPoint .potx, looked up by name (never by index).
LAYOUT_CANDIDATES: dict[SlideLayout, tuple[str, ...]] = {
    SlideLayout.TITLE: ("2_Title Slide", "Title Slide"),
    SlideLayout.BULLETS: ("Title and Content",),
    SlideLayout.TABLE: ("Title and Content", "Title Only"),
    SlideLayout.BULLETS_AND_TABLE: ("Two Content",),
}

_SKIP_PLACEHOLDERS = {
    PP_PLACEHOLDER.FOOTER,
    PP_PLACEHOLDER.SLIDE_NUMBER,
    PP_PLACEHOLDER.DATE,
    PP_PLACEHOLDER.HEADER,
}
_TITLE_PLACEHOLDERS = {
    PP_PLACEHOLDER.TITLE,
    PP_PLACEHOLDER.CENTER_TITLE,
    PP_PLACEHOLDER.VERTICAL_TITLE,
}


def template_path() -> Path:
    """Company master: env override, packaged copy, then repo root."""
    override = os.environ.get("MKTG_DECK_TEMPLATE", "").strip()
    if override:
        return Path(override)
    packaged = Path(__file__).resolve().parent / "templates" / TEMPLATE_FILENAME
    if packaged.is_file():
        return packaged
    return Path(__file__).resolve().parents[3] / TEMPLATE_FILENAME


def _materialize_template(source: Path, dest_dir: Path) -> Path:
    """python-pptx rejects .potx content types; rewrite as a .pptx copy."""
    dest = dest_dir / "master.pptx"
    raw = source.read_bytes()
    with zipfile.ZipFile(io.BytesIO(raw)) as zin:
        content_types = zin.read("[Content_Types].xml").decode("utf-8")
        if TEMPLATE_CT not in content_types:
            dest.write_bytes(raw)
            return dest
        content_types = content_types.replace(TEMPLATE_CT, PRESENTATION_CT)
        with zipfile.ZipFile(dest, "w") as zout:
            for item in zin.infolist():
                data = (
                    content_types.encode("utf-8")
                    if item.filename == "[Content_Types].xml"
                    else zin.read(item.filename)
                )
                zout.writestr(item, data)
    return dest


def _strip_existing_slides(prs: Presentation) -> None:
    sld_id_lst = prs.slides._sldIdLst
    for sld_id in list(sld_id_lst):
        r_id = sld_id.get(qn("r:id"))
        if r_id:
            prs.part.drop_rel(r_id)
        sld_id_lst.remove(sld_id)


def _layout(prs: Presentation, kind: SlideLayout):
    names = LAYOUT_CANDIDATES[kind]
    available = {item.name: item for item in prs.slide_layouts}
    for name in names:
        if name in available:
            return available[name]
    raise ValueError(
        f"SailPoint template is missing layout {names[0]!r}. "
        f"Available: {', '.join(sorted(available))}"
    )


def _placeholder_type(shape) -> object | None:
    try:
        return shape.placeholder_format.type
    except (AttributeError, ValueError):
        return None


def _set_text(shape, text: str) -> None:
    if shape is None or not getattr(shape, "has_text_frame", False):
        return
    frame = shape.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.paragraphs[0].text = text or ""


def _fill_title(slide_obj, title: str, subtitle: str | None = None) -> None:
    if slide_obj.shapes.title is not None:
        _set_text(slide_obj.shapes.title, title)
    else:
        for shape in slide_obj.placeholders:
            if _placeholder_type(shape) in _TITLE_PLACEHOLDERS:
                _set_text(shape, title)
                break
    if subtitle is None:
        return
    for shape in slide_obj.placeholders:
        if _placeholder_type(shape) == PP_PLACEHOLDER.SUBTITLE:
            _set_text(shape, subtitle)
            return


def _content_placeholders(slide_obj) -> list:
    found = []
    for shape in slide_obj.placeholders:
        kind = _placeholder_type(shape)
        if kind in _SKIP_PLACEHOLDERS or kind in _TITLE_PLACEHOLDERS:
            continue
        if kind == PP_PLACEHOLDER.SUBTITLE:
            continue
        found.append(shape)
    return found


def _add_bullets(shape, bullets: list[str]) -> None:
    if shape is None or not getattr(shape, "has_text_frame", False):
        return
    lines = [item for item in bullets if item]
    if not lines:
        return
    frame = shape.text_frame
    frame.clear()
    frame.word_wrap = True
    for i, bullet in enumerate(lines):
        para = frame.paragraphs[0] if i == 0 else frame.add_paragraph()
        para.text = bullet
        para.level = 0


def _fill_table(slide_obj, placeholder, columns: list[str], rows: list[list[str]]) -> None:
    if not columns or placeholder is None:
        return
    rows = rows[:10]
    if placeholder.has_text_frame:
        placeholder.text_frame.clear()
    table = slide_obj.shapes.add_table(
        len(rows) + 1, len(columns),
        placeholder.left, placeholder.top,
        placeholder.width, placeholder.height,
    ).table
    for col, name in enumerate(columns):
        table.cell(0, col).text = str(name)
    for r, row in enumerate(rows, start=1):
        for c in range(len(columns)):
            table.cell(r, c).text = str(row[c]) if c < len(row) else ""


def _body_lines(slide: Slide) -> list[str]:
    lines = list(slide.bullets)
    if slide.takeaway:
        lines.append(slide.takeaway)
    return lines


def _add_title_slide(prs: Presentation, slide: Slide) -> None:
    s = prs.slides.add_slide(_layout(prs, SlideLayout.TITLE))
    _fill_title(s, slide.title, slide.subtitle)


def _add_content_slide(prs: Presentation, slide: Slide) -> None:
    kind = slide.layout
    s = prs.slides.add_slide(_layout(prs, kind))
    _fill_title(s, slide.title)
    bodies = _content_placeholders(s)

    if kind == SlideLayout.BULLETS:
        if bodies:
            _add_bullets(bodies[0], _body_lines(slide))
        return

    if kind == SlideLayout.TABLE:
        if bodies:
            _fill_table(s, bodies[0], slide.table_columns, slide.table_rows)
        if slide.takeaway and len(bodies) > 1:
            _add_bullets(bodies[1], [slide.takeaway])
        return

    if kind == SlideLayout.BULLETS_AND_TABLE:
        if bodies:
            _add_bullets(bodies[0], _body_lines(slide))
        if len(bodies) > 1:
            _fill_table(s, bodies[1], slide.table_columns, slide.table_rows)
        elif bodies:
            _fill_table(s, bodies[0], slide.table_columns, slide.table_rows)


def render_deck(
    spec: DeckSpec,
    filename: str,
    output_dir: Path | None = None,
    template: Path | None = None,
) -> Path:
    """Write the deck to disk and return where it went."""
    directory = output_dir or OUTPUT_DIR
    directory.mkdir(parents=True, exist_ok=True)
    source = template or template_path()
    if not source.is_file():
        raise FileNotFoundError(
            f"SailPoint deck template was not found at {source}"
        )

    with tempfile.TemporaryDirectory() as tmp:
        master = _materialize_template(source, Path(tmp))
        prs = Presentation(str(master))
        _strip_existing_slides(prs)

        _add_title_slide(prs, Slide(
            layout=SlideLayout.TITLE, title=spec.title, subtitle=spec.subtitle))

        for slide in spec.slides:
            if slide.layout == SlideLayout.TITLE:
                _add_title_slide(prs, slide)
            else:
                _add_content_slide(prs, slide)

        if not filename.endswith(".pptx"):
            filename += ".pptx"
        path = directory / filename
        prs.save(str(path))
        return path


# Keep Emu imported so tests can compare slide size against the master.
SAILPOINT_WIDTH = Emu(12192000)
SAILPOINT_HEIGHT = Emu(6858000)
