"""Tests for the tool layer.

Everything here runs without a model. That is the point of pushing the real
work into tools: the parts that can be wrong in a damaging way (validation
rules, filters, file output) are ordinary Python and can be tested ordinarily.
"""

from __future__ import annotations

import json
from types import SimpleNamespace

import pytest

from mktg_core.contracts import DeckSpec
from mktg_core.contracts.marketing_ops import UPLOADED_LEAD_FILES_STATE
from mktg_core.rendering.artifact import save_artifact
from tools.abm.abm_tools import get_gap_value_map, get_top_accounts
from tools.analysis.deck_tools import save_deck
from tools.brand.brand_tools import get_share_of_voice, get_social_sentiment
from tools.campaign_design.campaign_design_tools import (
    get_competitor_intelligence,
    get_named_campaign,
)
from tools.regional_events.events_tools import (
    find_accounts_in_city,
    get_account_snapshot,
    get_event_attendees,
    prepare_account_briefs,
)
from tools.marketing_ops.mops_tools import (
    build_6sense_segment,
    clean_lead_rows,
    create_sfdc_campaign,
    load_cleaned_list,
    read_lead_file,
    validate_lead_rows,
)


def fake_context():
    """The tools only ever touch `.state`, so a plain dict holder is enough."""
    return SimpleNamespace(state={})


# --- deck rendering -------------------------------------------------------

def test_save_deck_writes_a_real_file(tmp_path, monkeypatch):
    import mktg_core.rendering.deck as deck_module
    from mktg_core.rendering.deck import template_path
    from pptx import Presentation

    if not template_path().is_file():
        pytest.skip("SailPoint .potx is not on disk")
    monkeypatch.setattr(deck_module, "OUTPUT_DIR", tmp_path)

    spec = {
        "title": "Pipeline Review",
        "subtitle": "Q3 2026",
        "slides": [
            {"layout": "bullets", "title": "Position",
             "bullets": ["EMEA is short", "AMER is fine"],
             "takeaway": "EMEA is the gap"},
            {"layout": "table", "title": "Coverage",
             "table_columns": ["Region", "Coverage"],
             "table_rows": [["AMER", "2.39x"], ["EMEA", "1.46x"]]},
        ],
    }
    result = save_deck(json.dumps(spec), "test-deck")

    assert "saved to" in result.lower()
    assert "output/decks/test-deck.pptx" in result
    written = list(tmp_path.glob("*.pptx"))
    assert len(written) == 1
    assert written[0].stat().st_size > 1000
    prs = Presentation(str(written[0]))
    assert int(prs.slide_width) == int(deck_module.SAILPOINT_WIDTH)
    assert int(prs.slide_height) == int(deck_module.SAILPOINT_HEIGHT)
    names = [slide.slide_layout.name for slide in prs.slides]
    assert names[0] == "2_Title Slide"
    allowed = {"2_Title Slide", "Title Slide", "Title and Content", "Title Only", "Two Content"}
    assert set(names) <= allowed
    assert len(prs.slides) == 3


def test_save_deck_explains_broken_json_instead_of_crashing():
    result = save_deck("{not json at all", "x")
    assert "could not be parsed" in result
    assert "valid JSON" in result


def test_save_deck_explains_a_wrong_shape_and_shows_the_expected_one():
    """The agent has to be able to correct itself from the error text."""
    result = save_deck(json.dumps({"slides": [{"layout": "nope", "title": "x"}]}), "x")
    assert "did not match" in result
    assert "Valid layouts" in result


def test_save_deck_refuses_an_empty_deck():
    result = save_deck(json.dumps({"title": "Empty", "slides": []}), "x")
    assert "no slides" in result


def test_every_layout_renders(tmp_path, monkeypatch):
    import mktg_core.rendering.deck as deck_module
    from mktg_core.rendering.deck import template_path
    from pptx import Presentation

    if not template_path().is_file():
        pytest.skip("SailPoint .potx is not on disk")
    monkeypatch.setattr(deck_module, "OUTPUT_DIR", tmp_path)

    spec = DeckSpec.model_validate({
        "title": "All layouts",
        "slides": [
            {"layout": "title", "title": "A section", "subtitle": "sub"},
            {"layout": "bullets", "title": "B", "bullets": ["one", "two"]},
            {"layout": "table", "title": "C", "table_columns": ["x"],
             "table_rows": [["1"]]},
            {"layout": "bullets_and_table", "title": "D", "bullets": ["note"],
             "table_columns": ["x", "y"], "table_rows": [["1", "2"]],
             "takeaway": "done"},
        ],
    })
    result = save_deck(spec.model_dump_json(), "layouts")
    assert "saved to" in result.lower()
    prs = Presentation(str(next(tmp_path.glob("*.pptx"))))
    names = [slide.slide_layout.name for slide in prs.slides]
    assert "Two Content" in names
    assert "Title and Content" in names


def test_missing_sailpoint_layout_name_fails_clearly(tmp_path, monkeypatch):
    import mktg_core.rendering.deck as deck_module
    from mktg_core.contracts import SlideLayout
    from mktg_core.rendering.deck import template_path

    if not template_path().is_file():
        pytest.skip("SailPoint .potx is not on disk")
    monkeypatch.setattr(deck_module, "OUTPUT_DIR", tmp_path)
    monkeypatch.setattr(deck_module, "LAYOUT_CANDIDATES", {
        **deck_module.LAYOUT_CANDIDATES,
        SlideLayout.BULLETS: ("Not A Real Layout",),
    })
    spec = {
        "title": "X",
        "slides": [{"layout": "bullets", "title": "B", "bullets": ["one"]}],
    }
    with pytest.raises(ValueError, match="missing layout 'Not A Real Layout'"):
        save_deck(json.dumps(spec), "bad")


def test_long_tables_are_capped_rather_than_overflowing(tmp_path, monkeypatch):
    import mktg_core.rendering.deck as deck_module
    from mktg_core.rendering.deck import template_path

    if not template_path().is_file():
        pytest.skip("SailPoint .potx is not on disk")
    monkeypatch.setattr(deck_module, "OUTPUT_DIR", tmp_path)

    spec = {"title": "Big", "slides": [{
        "layout": "table", "title": "Many rows",
        "table_columns": ["n"],
        "table_rows": [[str(i)] for i in range(50)],
    }]}
    assert "saved to" in save_deck(json.dumps(spec), "big")


# --- Salesforce campaign creation ----------------------------------------

def test_create_campaign_returns_an_id():
    result = create_sfdc_campaign(
        "EMEA-Webinar-2026Q4-Zero-Trust", "Webinar", "EMEA",
        "2026-10-01", "2026-12-15", 25000, "owner@gru.com", "Zero Trust webinar")
    assert "## Receipt" in result
    assert "Id: 701MOCK" in result
    assert "MOCK" in result, "the answer must make clear nothing real was written"
    assert "Nothing was written to a real Salesforce org" in result
    assert "Type: Webinar" in result
    assert "Dates: 2026-10-01 to 2026-12-15" in result


def test_create_campaign_warns_about_a_non_conforming_name():
    result = create_sfdc_campaign(
        "some webinar", "Webinar", "EMEA", "2026-10-01", "2026-12-15",
        25000, "owner@gru.com", "")
    assert "does not match the convention" in result
    assert "REGION-TYPE-YYYYQN" in result


def test_create_campaign_rejects_an_unknown_type_and_lists_the_valid_ones():
    result = create_sfdc_campaign(
        "EMEA-Telepathy-2026Q4-X", "Telepathy", "EMEA", "2026-10-01",
        "2026-12-15", 1000, "o@gru.com", "")
    assert "Could not create" in result
    assert "Webinar" in result, "the error should tell the agent what is valid"


def test_create_campaign_rejects_an_end_date_before_the_start():
    result = create_sfdc_campaign(
        "EMEA-Webinar-2026Q4-X", "Webinar", "EMEA", "2026-12-01",
        "2026-10-01", 1000, "o@gru.com", "")
    assert "before start date" in result


def test_create_campaign_is_reproducible():
    """Same request, same id. Reproducible demos matter more than realism here."""
    args = ("EMEA-Webinar-2026Q4-Repeat", "Webinar", "EMEA", "2026-10-01",
            "2026-12-15", 1000, "o@gru.com", "")
    assert create_sfdc_campaign(*args) == create_sfdc_campaign(*args)


# --- 6sense segments ------------------------------------------------------

def test_segment_filters_actually_narrow_the_result():
    broad = build_6sense_segment("broad", "", "", "", "", 0, 0, False)
    narrow = build_6sense_segment(
        "narrow", "Financial Services", "EMEA", "", "Zero Trust", 80, 0, True)

    broad_count = int(broad.split("Accounts matched: ")[1].split("\n")[0])
    narrow_count = int(narrow.split("Accounts matched: ")[1].split("\n")[0])
    assert 0 < narrow_count < broad_count


def test_segment_echoes_the_filters_back():
    """The user cannot catch a wrong assumption they were never shown."""
    result = build_6sense_segment(
        "zt", "Financial Services", "EMEA", "", "Zero Trust", 80, 0, False)
    assert "Financial Services" in result
    assert "intent score >= 80" in result
    assert "SEG-MOCK-" in result


def test_empty_segment_suggests_how_to_fix_it():
    result = build_6sense_segment(
        "impossible", "Retail", "APJ", "", "Zero Trust", 99, 100000, True)
    assert "0 accounts" in result
    assert "too narrow" in result


def test_segment_rejects_an_unknown_region():
    assert "Unknown region" in build_6sense_segment(
        "x", "", "ATLANTIS", "", "", 0, 0, False)


# --- the list load chain --------------------------------------------------

def test_read_step_reports_headers_and_stores_the_rows():
    ctx = fake_context()
    result = read_lead_file("event_leads_munich.csv", ctx)
    assert "All required columns are present" in result
    assert len(ctx.state["mops_list_load_raw_rows"]) == 10


def test_read_step_uses_rows_already_in_session():
    ctx = fake_context()
    ctx.state[UPLOADED_LEAD_FILES_STATE] = {
        "uploaded.csv": [
            {
                "First Name": "Ada",
                "Last Name": "Lovelace",
                "Email": "ada@example.com",
                "Company": "Vantage",
                "Country": "DE",
            }
        ]
    }
    result = read_lead_file("uploaded.csv", ctx)
    assert "All required columns are present" in result
    assert ctx.state["mops_list_load_raw_rows"][0]["Email"] == "ada@example.com"


def test_read_step_handles_a_missing_file_gracefully():
    assert "No file named" in read_lead_file("nope.csv", fake_context())


def test_clean_step_expands_country_codes_and_lowercases_emails():
    ctx = fake_context()
    read_lead_file("event_leads_munich.csv", ctx)
    result = clean_lead_rows(ctx)

    assert "DE -> Germany" in result
    assert "lowercased email" in result

    cleaned = ctx.state["mops_list_load_clean_rows"]
    assert all(row["Country"] not in {"DE", "FR", "IT", "NL", "AT"}
               for row in cleaned)
    assert all(row["Email"] == row["Email"].lower() for row in cleaned)
    assert all(v == v.strip() for row in cleaned for v in row.values())


def test_clean_step_rejects_nothing():
    """Cleaning and validating are separate steps on purpose."""
    ctx = fake_context()
    read_lead_file("event_leads_munich.csv", ctx)
    clean_lead_rows(ctx)
    assert len(ctx.state["mops_list_load_clean_rows"]) == 10


def test_validate_step_catches_every_planted_problem():
    ctx = fake_context()
    read_lead_file("event_leads_munich.csv", ctx)
    clean_lead_rows(ctx)
    result = validate_lead_rows(ctx)

    assert "required field is empty" in result       # blank email
    assert "not a valid email" in result             # tom.bennett[at]...
    assert "duplicate" in result                     # repeated Anna Keller
    assert "suppressed" in result                    # gru-internal.com

    valid = ctx.state["mops_list_load_valid_rows"]
    assert 0 < len(valid) < 10

    emails = [r["Email"] for r in valid]
    assert len(emails) == len(set(emails)), "duplicates must not survive"
    assert not any("gru-internal.com" in e for e in emails)


def test_validate_step_keeps_the_good_rows():
    ctx = fake_context()
    read_lead_file("event_leads_munich.csv", ctx)
    clean_lead_rows(ctx)
    validate_lead_rows(ctx)
    emails = [r["Email"] for r in ctx.state["mops_list_load_valid_rows"]]
    assert "anna.keller@northwindcapital.com" in emails
    assert "hannah.weber@solsticecapital.com" in emails, "was uppercase in source"


def test_load_step_only_uploads_validated_rows():
    ctx = fake_context()
    read_lead_file("event_leads_munich.csv", ctx)
    clean_lead_rows(ctx)
    validate_lead_rows(ctx)
    expected = len(ctx.state["mops_list_load_valid_rows"])

    result = load_cleaned_list("Munich Event Leads", ctx)
    assert f"Rows loaded: {expected}" in result
    assert "Germany" in result


@pytest.mark.parametrize("step,expected", [
    (clean_lead_rows, "Run read_lead_file first"),
    (validate_lead_rows, "Run clean_lead_rows first"),
])
def test_steps_refuse_to_run_out_of_order(step, expected):
    """Belt and braces. The SequentialAgent enforces order, and so does this."""
    assert expected in step(fake_context())


def test_load_refuses_when_nothing_was_validated():
    assert "Run validate_lead_rows first" in load_cleaned_list("x", fake_context())


# --- events tools ---------------------------------------------------------

def test_account_snapshot_covers_a_prospect_and_a_customer():
    from mktg_core.connectors import get_connectors
    c = get_connectors()
    with_opp = {o.account_id for o in c.sfdc.list_opportunities()}
    accounts = c.sfdc.list_accounts()

    prospect = next(a for a in accounts if a.id not in with_opp)
    customer = next(a for a in accounts if a.id in with_opp)

    assert "unworked account" in get_account_snapshot(prospect.id)
    assert "Open pipeline" in get_account_snapshot(customer.id)


def test_account_snapshot_handles_an_unknown_id():
    assert "No account found" in get_account_snapshot("ACC-999999")


def test_city_lookup_lists_known_cities_when_it_misses():
    result = find_accounts_in_city("Atlantis")
    assert "No accounts found" in result
    assert "Munich" in result, "the error should help the agent retry"


def test_attendee_list_ends_with_ids_the_next_tool_can_use():
    """The tools have to chain, and the model reads this line to do it."""
    result = get_event_attendees("EVT-001")
    assert "Account ids in the room:" in result
    ids = result.split("Account ids in the room:")[1].strip()
    assert ids.startswith("ACC-")


def test_brief_preparation_splits_work_across_the_groups():
    ctx = fake_context()
    result = prepare_account_briefs("", "EVT-001", ctx)

    assert "Prepared briefing data" in result
    groups = [ctx.state.get(f"events_brief_data_{n}", "") for n in (1, 2, 3)]
    assert all(groups), "every group should get accounts for a room this size"
    assert sum(g.count("### ") for g in groups) >= 8


def test_brief_preparation_accepts_explicit_account_ids():
    ctx = fake_context()
    result = prepare_account_briefs("ACC-001,ACC-002,ACC-003", "", ctx)
    assert "3 accounts" in result


def test_brief_preparation_asks_for_input_when_given_none():
    assert "Provide either" in prepare_account_briefs("", "", fake_context())


def test_brief_preparation_ignores_unknown_ids_but_still_works():
    ctx = fake_context()
    result = prepare_account_briefs("ACC-001,ACC-999999", "", ctx)
    assert "1 accounts" in result
    assert "Ignored 1 unknown" in result


def test_abm_top_accounts_returns_tiers():
    result = get_top_accounts("EMEA")
    assert "Tier" in result
    assert "Composite score" in result


def test_gap_value_map_quantifies_roi():
    from mktg_core.connectors import get_connectors
    munich = next(
        a for a in get_connectors().sfdc.list_accounts()
        if a.city == "Munich" and a.is_target_account
    )
    result = get_gap_value_map(munich.id)
    assert "Annual value (USD)" in result
    assert "Total quantified" in result


def test_named_campaign_lookup_finds_paid_social():
    result = get_named_campaign("Cloud Security Always On")
    assert "CMP-001" in result
    assert "Paid Social" in result
    assert "Spend (USD): 86000" in result


def test_named_campaign_lookup_explains_unknown_names():
    result = get_named_campaign("Not A Real Campaign")
    assert "No campaign matched" in result
    assert "Cloud Security Always On" in result


def test_brand_sentiment_and_sov():
    assert "Positive" in get_social_sentiment()
    assert "SailPoint" in get_share_of_voice()
    assert "CyberArk" in get_share_of_voice()


def test_save_battlecard_artifact(tmp_path, monkeypatch):
    import mktg_core.rendering.artifact as art
    monkeypatch.setattr(art, "OUTPUT_DIR", tmp_path)

    spec = {
        "competitor": "CyberArk",
        "positioning_statement": "Governance beyond vaulting.",
        "why_we_win": ["Deep IGA"],
        "their_strengths": ["PAM"],
        "landmines": [],
        "objection_handlers": [],
        "displacement_cta": "Book a review",
    }
    result = save_artifact(json.dumps(spec), "battlecard", "cyberark.json")
    assert result.name == "cyberark.md"
    assert "CyberArk" in result.read_text()
    from mktg_core.rendering.artifact import describe_saved

    described = describe_saved(result)
    assert described.startswith("Saved to ")
    assert "Governance beyond vaulting" in described
